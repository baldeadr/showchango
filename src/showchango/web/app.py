from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import tempfile
import unicodedata
from pathlib import Path

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from showchango.analytics.librosa_analyzer import LibrosaAnalyzer
from showchango.core.curva import datos_para_chartjs
from showchango.core.diagnostico import diagnosticar, listar_plantillas
from showchango.core.esquema import (
    OrdenPayload,
    Pista,
    Proyecto,
    a_texto,
    de_texto,
    nuevo_id,
    nuevo_proyecto,
)
from showchango.db import Cache
from showchango.web.sesiones import Sesiones

PLANTILLAS = Path(__file__).parent / "plantillas"
STATIC = Path(__file__).parent / "static"
COOKIE_SESION = "sc_sesion"
EXTENSIONES_AUDIO = {".mp3", ".wav", ".flac", ".ogg", ".m4a", ".aac", ".wma"}
TAMANO_MAXIMO_AUDIO = 50 * 1024 * 1024  # 50 MB
_ARCHIVO_REQUERIDO = File(...)

_MAPEO_CSV = {
    "titulo": ["titulo", "titulo", "title", "nombre", "track", "cancion", "cancion"],
    "artista": ["artista", "artista", "artist", "autor", "banda", "dj"],
    "bpm": ["bpm", "tempo"],
    "energia": ["energia", "energia", "energy"],
    "bailabilidad": ["bailabilidad", "danceability", "dance"],
    "duracion_s": [
        "duracion",
        "duracion",
        "duracion_s",
        "duracions",
        "duration",
        "length",
        "dur",
        "tiempo",
    ],
    "loudness_db": ["loudness", "loudness_db", "volumen", "sonoridad"],
    "valencia": ["valencia", "valence", "mood", "animo", "animo"],
    "voz": ["voz", "voice", "vocals", "vocal"],
    "genero": ["genero", "genero", "genre", "estilo"],
    "tonalidad": ["tonalidad", "key", "tono"],
    "notas": ["notas", "notes", "comentarios", "observaciones"],
}


def _normalizar(texto: str) -> str:
    texto = texto.strip().lower()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-z0-9_]", "_", texto)


def _campo_desde_encabezado(encabezado: str) -> str | None:
    clave = _normalizar(encabezado)
    for campo, sinonimos in _MAPEO_CSV.items():
        if clave in sinonimos or clave.replace("_", "") in [s.replace("_", "") for s in sinonimos]:
            return campo
    return None


def _float_or_none(valor: str) -> float | None:
    if valor is None:
        return None
    valor = str(valor).strip().replace(",", ".")
    if valor == "":
        return None
    if ":" in valor:
        partes = valor.split(":")
        msg_duracion = (
            f"Formato de duración no válido: '{valor}'. "
            "Usa segundos (190) o minutos:segundos (3:10)."
        )
        if len(partes) != 2:
            raise ValueError(msg_duracion)
        try:
            minutos = float(partes[0])
            segundos = float(partes[1])
        except ValueError as exc:
            raise ValueError(msg_duracion) from exc
        if minutos < 0 or segundos < 0 or segundos >= 60:
            raise ValueError(f"Duración no válida: '{valor}'.")
        return round(minutos * 60 + segundos, 6)
    try:
        return float(valor)
    except ValueError as exc:
        raise ValueError(f"Valor numérico no válido: '{valor}'.") from exc


def _voz_or_none(valor: str) -> str | None:
    if not valor:
        return None
    valor = valor.strip().lower()
    if valor in {"vocal", "voz"}:
        return "vocal"
    if valor in {"instrumental", "inst"}:
        return "instrumental"
    if valor in {"mixto", "ambos"}:
        return "mixto"
    return None


def _hash_archivo(ruta: Path) -> str:
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _titulo_desde_archivo(nombre: str | None) -> str:
    if not nombre:
        return "Audio importado"
    return Path(nombre).stem.replace("_", " ").replace("-", " ").strip().title()


def _crear_pista_desde_features(features: dict, titulo: str) -> Pista:
    return Pista(
        id=nuevo_id(),
        titulo=titulo,
        fuente="analisis",
        bpm=features.get("bpm"),
        energia=features.get("energia"),
        bailabilidad=features.get("bailabilidad"),
        duracion_s=features.get("duracion_s"),
        loudness_db=features.get("loudness_db"),
        valencia=features.get("valencia"),
        voz=features.get("voz"),
        genero=features.get("genero"),
        tonalidad=features.get("tonalidad"),
        notas=features.get("notas", ""),
    )


def _fila_a_pista(fila: dict) -> Pista:
    datos: dict[str, str | None] = {campo: None for campo in _MAPEO_CSV}
    for encabezado, valor in fila.items():
        campo = _campo_desde_encabezado(encabezado)
        if campo:
            datos[campo] = str(valor) if valor is not None else None

    return Pista(
        id=nuevo_id(),
        titulo=datos.get("titulo") or "Sin título",
        artista=datos.get("artista") or "",
        bpm=_float_or_none(datos.get("bpm")),
        energia=_float_or_none(datos.get("energia")),
        bailabilidad=_float_or_none(datos.get("bailabilidad")),
        duracion_s=_float_or_none(datos.get("duracion_s")),
        loudness_db=_float_or_none(datos.get("loudness_db")),
        valencia=_float_or_none(datos.get("valencia")),
        voz=_voz_or_none(datos.get("voz") or ""),
        genero=datos.get("genero") or None,
        tonalidad=datos.get("tonalidad") or None,
        fuente="csv",
        notas=datos.get("notas") or "",
    )


def _parse_csv(contenido: bytes) -> list[Pista]:
    texto = contenido.decode("utf-8-sig")
    lector = csv.DictReader(io.StringIO(texto))
    return [_fila_a_pista(fila) for fila in lector]


def _parse_excel(contenido: bytes) -> list[Pista]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = wb.active
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return []
    encabezados = [str(c) if c is not None else "" for c in filas[0]]
    pistas: list[Pista] = []
    for fila in filas[1:]:
        datos = {encabezados[i]: fila[i] for i in range(min(len(encabezados), len(fila)))}
        pistas.append(_fila_a_pista(datos))
    return pistas


def crear_app() -> FastAPI:
    app = FastAPI(title="Show Chango")
    app.mount("/static", StaticFiles(directory=STATIC), name="static")
    templates = Jinja2Templates(directory=PLANTILLAS)
    sesiones = Sesiones()
    cache = Cache()
    analyzer = LibrosaAnalyzer()
    jobs: dict[str, dict] = {}

    def _sid(request: Request) -> str | None:
        return request.cookies.get(COOKIE_SESION)

    def _respuesta_con_sesion(sid: str, url: str) -> RedirectResponse:
        respuesta = RedirectResponse(url, status_code=303)
        respuesta.set_cookie(COOKIE_SESION, sid, max_age=8 * 3600, httponly=True)
        return respuesta

    def _proyecto_o_redir(sid: str | None) -> Proyecto | RedirectResponse:
        if sid is None:
            return RedirectResponse("/", status_code=303)
        proyecto = sesiones.obtener_proyecto(sid)
        if proyecto is None:
            return RedirectResponse("/", status_code=303)
        return proyecto

    def _render_set(
        request: Request,
        proyecto: Proyecto,
        error: str | None = None,
        plantilla_id: str = "climax_70",
    ):
        plantillas_validas = {"climax_70", "construccion_dj", "picos_rock"}
        plantilla_id = plantilla_id if plantilla_id in plantillas_validas else "climax_70"
        diagnostico = diagnosticar(proyecto, plantilla_id) if proyecto.orden else None
        return templates.TemplateResponse(
            request,
            "set.html",
            {
                "proyecto": proyecto,
                "curva_json": json.dumps(
                    datos_para_chartjs(proyecto), ensure_ascii=False
                ),
                "error": error,
                "diagnostico": diagnostico,
                "plantilla_id": plantilla_id,
                "plantillas": listar_plantillas(),
            },
        )

    def _analizar_audio(
        job_id: str,
        ruta: Path,
        sid: str,
        titulo: str,
        hash_archivo: str,
    ) -> None:
        try:
            jobs[job_id]["status"] = "analyzing"
            features = analyzer.analyze(ruta)
            cache.guardar(hash_archivo, features, analyzer.version())
            pista = _crear_pista_desde_features(features, titulo)
            proyecto = sesiones.obtener_proyecto(sid)
            if proyecto is not None:
                proyecto.pistas.append(pista)
                proyecto.orden.append(pista.id)
                sesiones.guardar_proyecto(sid, proyecto)
            jobs[job_id]["status"] = "done"
        except Exception as exc:
            jobs[job_id]["status"] = "error"
            jobs[job_id]["message"] = str(exc)
        finally:
            try:
                os.remove(ruta)
            except OSError:
                pass

    @app.get("/", response_class=HTMLResponse)
    def home(request: Request):
        sid = _sid(request)
        proyecto = sesiones.obtener_proyecto(sid) if sid else None
        return templates.TemplateResponse(
            request,
            "index.html",
            {"proyecto": proyecto},
        )

    @app.post("/set/nuevo")
    def crear_set(
        request: Request,
        nombre: str = Form(""),
        artista: str = Form(""),
    ):
        sid = _sid(request) or sesiones.nueva()
        proyecto = nuevo_proyecto(nombre=nombre, artista=artista)
        sesiones.guardar_proyecto(sid, proyecto)
        return _respuesta_con_sesion(sid, "/set")

    @app.get("/set", response_class=HTMLResponse)
    def ver_set(request: Request, plantilla: str = Query("climax_70")):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        return _render_set(request, resultado, plantilla_id=plantilla)

    @app.get("/set/curva-json")
    def curva_json(request: Request):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        return datos_para_chartjs(resultado)

    @app.post("/set/reordenar")
    async def reordenar(request: Request, payload: OrdenPayload):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        ids_validos = {p.id for p in resultado.pistas}
        orden = [pid for pid in payload.orden if pid in ids_validos]
        if set(orden) != ids_validos or len(orden) != len(ids_validos):
            raise HTTPException(status_code=400, detail="Orden inválido.")
        resultado.orden = orden
        sesiones.guardar_proyecto(sid, resultado)
        return datos_para_chartjs(resultado)

    @app.post("/set/importar")
    async def importar_set(request: Request, archivo: UploadFile = _ARCHIVO_REQUERIDO):
        sid = _sid(request)
        if sid is None:
            return RedirectResponse("/", status_code=303)
        contenido = await archivo.read()
        try:
            texto = contenido.decode("utf-8")
            proyecto = de_texto(texto)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"JSON inválido: {exc}") from exc
        sesiones.guardar_proyecto(sid, proyecto)
        return _respuesta_con_sesion(sid, "/set")

    @app.get("/set/exportar")
    def exportar_set(request: Request):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        contenido = a_texto(resultado)
        return Response(
            content=contenido,
            media_type="application/json",
            headers={
                "Content-Disposition": 'attachment; filename="showchango-set.json"',
            },
        )

    @app.post("/set/cerrar")
    def cerrar_set(request: Request):
        sid = _sid(request)
        if sid:
            sesiones.limpiar_proyecto(sid)
        respuesta = RedirectResponse("/", status_code=303)
        respuesta.delete_cookie(COOKIE_SESION)
        return respuesta

    @app.post("/pista/nueva")
    def crear_pista(
        request: Request,
        titulo: str = Form(""),
        artista: str = Form(""),
        bpm: str = Form(""),
        energia: str = Form(""),
        bailabilidad: str = Form(""),
        duracion_s: str = Form(""),
        loudness_db: str = Form(""),
        valencia: str = Form(""),
        voz: str = Form(""),
        genero: str = Form(""),
        tonalidad: str = Form(""),
        notas: str = Form(""),
    ):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        if not titulo.strip():
            return _render_set(request, resultado, error="El título es obligatorio.")
        try:
            pista = Pista(
                id=nuevo_id(),
                titulo=titulo.strip(),
                artista=artista.strip(),
                bpm=_float_or_none(bpm),
                energia=_float_or_none(energia),
                bailabilidad=_float_or_none(bailabilidad),
                duracion_s=_float_or_none(duracion_s),
                loudness_db=_float_or_none(loudness_db),
                valencia=_float_or_none(valencia),
                voz=_voz_or_none(voz),
                genero=genero.strip() or None,
                tonalidad=tonalidad.strip() or None,
                fuente="manual",
                notas=notas.strip(),
            )
        except ValueError as exc:
            return _render_set(request, resultado, error=str(exc))
        resultado.pistas.append(pista)
        resultado.orden.append(pista.id)
        sesiones.guardar_proyecto(sid, resultado)
        return RedirectResponse("/set", status_code=303)

    @app.post("/pista/{pista_id}/actualizar")
    def actualizar_pista(
        request: Request,
        pista_id: str,
        titulo: str = Form(""),
        artista: str = Form(""),
        bpm: str = Form(""),
        energia: str = Form(""),
        bailabilidad: str = Form(""),
        duracion_s: str = Form(""),
        loudness_db: str = Form(""),
        valencia: str = Form(""),
        voz: str = Form(""),
        genero: str = Form(""),
        tonalidad: str = Form(""),
        notas: str = Form(""),
    ):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        for pista in resultado.pistas:
            if pista.id == pista_id:
                try:
                    pista.titulo = titulo.strip() or pista.titulo
                    pista.artista = artista.strip()
                    pista.bpm = _float_or_none(bpm)
                    pista.energia = _float_or_none(energia)
                    pista.bailabilidad = _float_or_none(bailabilidad)
                    pista.duracion_s = _float_or_none(duracion_s)
                    pista.loudness_db = _float_or_none(loudness_db)
                    pista.valencia = _float_or_none(valencia)
                    pista.voz = _voz_or_none(voz)
                    pista.genero = genero.strip() or None
                    pista.tonalidad = tonalidad.strip() or None
                    pista.notas = notas.strip()
                except ValueError as exc:
                    return _render_set(request, resultado, error=str(exc))
                break
        sesiones.guardar_proyecto(sid, resultado)
        return RedirectResponse("/set", status_code=303)

    @app.post("/pista/{pista_id}/eliminar")
    def eliminar_pista(request: Request, pista_id: str):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        resultado.pistas = [p for p in resultado.pistas if p.id != pista_id]
        resultado.orden = [pid for pid in resultado.orden if pid != pista_id]
        resultado.transiciones = [t for t in resultado.transiciones if t.despues_de != pista_id]
        sesiones.guardar_proyecto(sid, resultado)
        return RedirectResponse("/set", status_code=303)

    @app.post("/set/importar-csv")
    async def importar_csv(request: Request, archivo: UploadFile = _ARCHIVO_REQUERIDO):
        sid = _sid(request)
        resultado = _proyecto_o_redir(sid)
        if isinstance(resultado, RedirectResponse):
            return resultado
        contenido = await archivo.read()
        nombre = archivo.filename or ""
        try:
            if nombre.lower().endswith((".xlsx", ".xls")):
                pistas = _parse_excel(contenido)
            else:
                pistas = _parse_csv(contenido)
        except Exception as exc:
            raise HTTPException(
                status_code=400, detail=f"No se pudo leer el archivo: {exc}"
            ) from exc
        for pista in pistas:
            if pista.titulo.strip():
                resultado.pistas.append(pista)
                resultado.orden.append(pista.id)
        sesiones.guardar_proyecto(sid, resultado)
        return RedirectResponse("/set", status_code=303)

    @app.post("/pista/subir-audio")
    async def subir_audio(
        request: Request,
        background_tasks: BackgroundTasks,
        archivo: UploadFile = _ARCHIVO_REQUERIDO,
    ):
        sid = _sid(request)
        if sid is None:
            return RedirectResponse("/", status_code=303)

        filename = archivo.filename or "audio"
        extension = Path(filename).suffix.lower()
        if extension not in EXTENSIONES_AUDIO:
            raise HTTPException(
                status_code=400,
                detail=f"Formato no soportado: {extension}. Usa {EXTENSIONES_AUDIO}.",
            )

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=extension)
        h = hashlib.sha256()
        tamano = 0
        while True:
            chunk = await archivo.read(8192)
            if not chunk:
                break
            tamano += len(chunk)
            if tamano > TAMANO_MAXIMO_AUDIO:
                tmp.close()
                os.remove(tmp.name)
                raise HTTPException(
                    status_code=413, detail="Archivo demasiado grande (máx. 50 MB)."
                )
            h.update(chunk)
            tmp.write(chunk)
        tmp.close()
        hash_archivo = h.hexdigest()
        titulo = _titulo_desde_archivo(filename)

        features = cache.obtener(hash_archivo)
        if features is not None:
            os.remove(tmp.name)
            proyecto = sesiones.obtener_proyecto(sid)
            if proyecto is None:
                return RedirectResponse("/", status_code=303)
            pista = _crear_pista_desde_features(features, titulo)
            proyecto.pistas.append(pista)
            proyecto.orden.append(pista.id)
            sesiones.guardar_proyecto(sid, proyecto)
            return RedirectResponse("/set", status_code=303)

        job_id = nuevo_id()
        jobs[job_id] = {"status": "pending", "message": "", "sid": sid}
        background_tasks.add_task(
            _analizar_audio, job_id, Path(tmp.name), sid, titulo, hash_archivo
        )
        return templates.TemplateResponse(
            request,
            "analisis_estado.html",
            {"job_id": job_id, "status": "pending", "message": ""},
        )

    @app.get("/pista/analisis/{job_id}/estado")
    def estado_analisis(request: Request, job_id: str):
        job = jobs.get(job_id)
        if job is None:
            return templates.TemplateResponse(
                request,
                "analisis_estado.html",
                {"job_id": None, "status": "error", "message": "Trabajo no encontrado."},
            )
        return templates.TemplateResponse(
            request,
            "analisis_estado.html",
            {"job_id": job_id, "status": job["status"], "message": job.get("message", "")},
        )

    return app


app = crear_app()
